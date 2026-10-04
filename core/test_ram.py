import json
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from core.admin import MemoriaRamAdmin, PlacaMadreAdmin
from core.compatibilidad import ACLARACION_RAM, MENSAJE_MEZCLA_RAM, LineaRam, validar_armado
from core.models import Carrito, ItemCarrito, MemoriaRam, PlacaMadre, Procesador, Proveedor
from core.tests import pieza


def ram(**cambios):
    datos = dict(
        tipo_ddr='DDR5',
        capacidad_gb=16,
        modulos_por_producto=1,
        capacidad_modulo_gb=16,
        formato_ram='DIMM',
        velocidad_mhz=5600,
    )
    datos.update(cambios)
    return pieza(**datos)


def placa(**cambios):
    datos = dict(
        socket_cpu='AM5',
        tipo_ram_soportado='DDR5',
        formato='ATX',
        ranuras_ram=4,
        formato_ram_soportado='DIMM',
        capacidad_maxima_ram_gb=128,
    )
    datos.update(cambios)
    return pieza(**datos)


def armar(**extra):
    datos = dict(
        procesador=pieza(socket='AM5'),
        gabinete=pieza(formato_soporte='ATX', largo_max_gpu_mm=400),
        tarjeta_grafica=pieza(largo_mm=300),
    )
    datos.update(extra)
    return validar_armado(**datos)


class ReglasRamTests(TestCase):
    def test_el_kit_ocupa_sus_modulos_y_no_duplica_la_capacidad(self):
        kit = ram(capacidad_gb=32, modulos_por_producto=2, capacidad_modulo_gb=16)
        resultado = armar(
            placa_madre=placa(ranuras_ram=2, capacidad_maxima_ram_gb=32),
            memorias_ram=[LineaRam(kit, 1)],
        )

        self.assertEqual(resultado.estado, 'compatible')
        self.assertEqual(resultado.advertencias, [])
        self.assertFalse(resultado.bloquea_agregar)

    def test_las_ranuras_exactas_cumplen_y_una_de_mas_no(self):
        modulo = ram()
        exacto = armar(
            placa_madre=placa(ranuras_ram=2, capacidad_maxima_ram_gb=64),
            memorias_ram=[LineaRam(modulo, 2)],
        )
        exceso = validar_armado(
            placa_madre=placa(ranuras_ram=2, capacidad_maxima_ram_gb=64),
            memorias_ram=[LineaRam(modulo, 3)],
        )

        self.assertEqual(exacto.estado, 'compatible')
        self.assertEqual(exceso.estado, 'incompatible')
        self.assertIn('3 ranuras', ' '.join(exceso.motivos_de_rechazo))
        self.assertIn('tiene 2', ' '.join(exceso.motivos_de_rechazo))

    def test_la_capacidad_maxima_exacta_cumple_y_un_gb_de_mas_no(self):
        modulo = ram(capacidad_gb=16)
        exacto = armar(
            placa_madre=placa(capacidad_maxima_ram_gb=32),
            memorias_ram=[LineaRam(modulo, 2)],
        )
        exceso = validar_armado(
            placa_madre=placa(capacidad_maxima_ram_gb=31),
            memorias_ram=[LineaRam(modulo, 2)],
        )

        self.assertEqual(exacto.estado, 'compatible')
        self.assertEqual(exceso.estado, 'incompatible')
        self.assertIn('32 GB', ' '.join(exceso.motivos_de_rechazo))
        self.assertIn('31 GB', ' '.join(exceso.motivos_de_rechazo))
        self.assertNotIn('dual', ' '.join(exceso.motivos_de_rechazo).lower())

    def test_mezclar_ddr_o_formatos_se_rechaza(self):
        ddr = validar_armado(
            placa_madre=placa(),
            memorias_ram=[LineaRam(ram(tipo_ddr='DDR4'), 1)],
        )
        formatos = validar_armado(
            memorias_ram=[
                LineaRam(ram(formato_ram='DIMM'), 1),
                LineaRam(ram(formato_ram='SO-DIMM'), 1),
            ],
        )

        self.assertEqual(ddr.estado, 'incompatible')
        self.assertIn('DDR4', ' '.join(ddr.motivos_de_rechazo))
        self.assertEqual(formatos.estado, 'incompatible')
        self.assertIn('SO-DIMM', ' '.join(formatos.motivos_de_rechazo))

    def test_mezclar_productos_compatibles_solo_advierte(self):
        uno = ram()
        otro = ram()
        resultado = armar(
            placa_madre=placa(),
            memorias_ram=[LineaRam(uno, 1), LineaRam(otro, 1)],
        )

        self.assertEqual(resultado.estado, 'compatible')
        self.assertFalse(resultado.bloquea_agregar)
        self.assertEqual(resultado.advertencias, [MENSAJE_MEZCLA_RAM])
        self.assertNotIn('dual', MENSAJE_MEZCLA_RAM.lower())

    def test_el_mismo_producto_repetido_no_advierte(self):
        modulo = ram()
        resultado = armar(
            placa_madre=placa(),
            memorias_ram=[LineaRam(modulo, 1), LineaRam(modulo, 1)],
        )

        self.assertEqual(resultado.estado, 'compatible')
        self.assertEqual(resultado.advertencias, [])

    def test_un_dato_desconocido_no_se_declara_compatible(self):
        sin_modulos = validar_armado(
            placa_madre=placa(),
            memoria_ram=ram(modulos_por_producto=None),
        )
        sin_formato = validar_armado(
            placa_madre=placa(),
            memoria_ram=ram(formato_ram=None),
        )
        sin_maximo = validar_armado(
            placa_madre=placa(capacidad_maxima_ram_gb=None),
            memoria_ram=ram(),
        )
        inconsistente = validar_armado(
            placa_madre=placa(),
            memoria_ram=ram(capacidad_gb=32, modulos_por_producto=2, capacidad_modulo_gb=8),
        )

        for resultado in (sin_modulos, sin_formato, sin_maximo, inconsistente):
            self.assertEqual(resultado.estado, 'datos_insuficientes')
            self.assertTrue(resultado.bloquea_agregar)

    def test_sin_la_otra_pieza_la_seleccion_queda_incompleta(self):
        solo_ram = validar_armado(memoria_ram=ram())
        solo_placa = validar_armado(placa_madre=placa())

        self.assertEqual(solo_ram.estado, 'incompleto')
        self.assertEqual(solo_placa.estado, 'incompleto')
        self.assertFalse(solo_ram.bloquea_agregar)
        self.assertFalse(solo_placa.bloquea_agregar)

    def test_el_admin_muestra_los_campos(self):
        self.assertIn('modulos_por_producto', MemoriaRamAdmin.list_display)
        self.assertIn('capacidad_modulo_gb', MemoriaRamAdmin.fieldsets[1][1]['fields'])
        self.assertIn('formato_ram', MemoriaRamAdmin.fieldsets[1][1]['fields'])
        self.assertIn('capacidad_maxima_ram_gb', PlacaMadreAdmin.list_display)
        self.assertIn('formato_ram_soportado', PlacaMadreAdmin.fieldsets[1][1]['fields'])


class ArmadoRamTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = User.objects.create_user('ram', 'ram@example.com', 'clave-ram')
        cls.proveedor = Proveedor.objects.create(nombre='Marca RAM')
        cls.cpu = Procesador.objects.create(
            proveedor=cls.proveedor, nombre='CPU RAM', precio=Decimal('10000'), stock=4,
            socket='AM5', nucleos=6, frecuencia_base=Decimal('3.50'),
        )
        cls.placa = PlacaMadre.objects.create(
            proveedor=cls.proveedor, nombre='Placa DIMM', precio=Decimal('20000'), stock=4,
            socket_cpu='AM5', chipset='B650', formato='ATX', ranuras_ram=4,
            tipo_ram_soportado='DDR5', formato_ram_soportado='DIMM', capacidad_maxima_ram_gb=64,
        )
        cls.modulo = cls._ram('Módulo 16', 16, 1, 16, 'DIMM', 'DDR5', 4, Decimal('10000'))
        cls.kit = cls._ram('Kit 32', 32, 2, 16, 'DIMM', 'DDR5', 4, Decimal('30000'))
        cls.otro = cls._ram('Otro módulo', 16, 1, 16, 'DIMM', 'DDR5', 2, Decimal('12000'))
        cls.sodimm = cls._ram('SO-DIMM', 16, 1, 16, 'SO-DIMM', 'DDR5', 4, Decimal('11000'))
        cls.ddr4 = cls._ram('DDR4', 16, 1, 16, 'DIMM', 'DDR4', 4, Decimal('9000'))
        cls.sin_dato = cls._ram('RAM sin módulos', 16, None, None, None, 'DDR5', 4, Decimal('8000'))

    @classmethod
    def _ram(cls, nombre, capacidad, modulos, por_modulo, formato, ddr, stock, precio):
        return MemoriaRam.objects.create(
            proveedor=cls.proveedor, nombre=nombre, precio=precio, stock=stock,
            capacidad_gb=capacidad, tipo_ddr=ddr, velocidad_mhz=5600,
            modulos_por_producto=modulos, capacidad_modulo_gb=por_modulo, formato_ram=formato,
        )

    def setUp(self):
        self.client.force_login(self.user)

    def post(self, url, componentes, extra=None):
        payload = {'componentes': componentes}
        if extra:
            payload.update(extra)
        return self.client.post(url, data=json.dumps(payload), content_type='application/json')

    def test_varias_lineas_recomendaciones_y_exportacion(self):
        agregada = self.post(reverse('core:recomendar_armado'), [
            {'tipo': 'placa_madre', 'id': self.placa.id},
            {'tipo': 'memoria_ram', 'id': self.modulo.id, 'cantidad': 1},
        ], {'categoria': 'memoria_ram', 'modo': 'agregar'})
        por_nombre = {item['nombre']: item for item in agregada.json()['candidatos']}
        self.assertEqual(por_nombre['Kit 32']['estado'], 'compatible')
        self.assertIn(MENSAJE_MEZCLA_RAM, por_nombre['Kit 32']['advertencias'])
        self.assertEqual(por_nombre['SO-DIMM']['estado'], 'incompatible')
        self.assertEqual(por_nombre['DDR4']['estado'], 'incompatible')

        reemplazo = self.post(reverse('core:recomendar_armado'), [
            {'tipo': 'placa_madre', 'id': self.placa.id},
            {'tipo': 'memoria_ram', 'id': self.sodimm.id, 'cantidad': 2},
        ], {
            'categoria': 'memoria_ram',
            'modo': 'reemplazar',
            'reemplaza_id': self.sodimm.id,
        })
        por_reemplazo = {item['nombre']: item for item in reemplazo.json()['candidatos']}
        self.assertEqual(por_reemplazo['Módulo 16']['estado'], 'compatible')
        self.assertEqual(por_reemplazo['SO-DIMM']['estado'], 'incompatible')

        evaluacion = self.post(reverse('core:evaluar_armado'), [
            {'tipo': 'placa_madre', 'id': self.placa.id},
            {'tipo': 'memoria_ram', 'id': self.modulo.id, 'cantidad': 1},
            {'tipo': 'memoria_ram', 'id': self.kit.id, 'cantidad': 1},
        ]).json()
        self.assertEqual(evaluacion['piezas']['memoria_ram']['estado'], 'compatible')
        self.assertEqual(evaluacion['estado'], 'incompleto')
        self.assertIn('3 de 4 ranuras', evaluacion['piezas']['memoria_ram']['etiqueta'])
        self.assertIn('48 GB de 64 GB', evaluacion['piezas']['memoria_ram']['etiqueta'])
        self.assertIn(MENSAJE_MEZCLA_RAM, evaluacion['advertencias'])
        self.assertNotIn('dual', evaluacion['piezas']['memoria_ram']['etiqueta'].lower())

        exportacion = self.post(reverse('core:exportar_armado'), [
            {'tipo': 'memoria_ram', 'id': self.modulo.id, 'cantidad': 2},
            {'tipo': 'memoria_ram', 'id': self.kit.id, 'cantidad': 1},
            {'tipo': 'placa_madre', 'id': self.placa.id},
        ])
        self.assertEqual(exportacion.status_code, 200)
        libro = load_workbook(BytesIO(exportacion.content))
        presupuesto = libro['Presupuesto']
        especificaciones = libro['Especificaciones']
        filas = {
            presupuesto.cell(fila, 2).value: fila
            for fila in range(5, 9)
            if presupuesto.cell(fila, 2).value
        }
        self.assertEqual(presupuesto.cell(filas['Módulo 16'], 3).value, 2)
        self.assertEqual(presupuesto.cell(filas['Módulo 16'], 5).value, 20000)
        self.assertEqual(presupuesto.cell(filas['Kit 32'], 3).value, 1)
        self.assertEqual(presupuesto.cell(filas['Kit 32'], 5).value, 30000)
        textos = ' '.join(
            str(valor)
            for hoja in (presupuesto, especificaciones)
            for fila in hoja.iter_rows(values_only=True)
            for valor in fila if valor
        )
        self.assertIn(ACLARACION_RAM, textos)
        self.assertIn(MENSAJE_MEZCLA_RAM, textos)
        self.assertIn('Capacidad del producto (GB)', textos)
        self.assertIn('Capacidad por módulo (GB)', textos)
        self.assertIn('Módulos por producto', textos)

    def test_el_stock_insuficiente_no_agrega_nada(self):
        carrito = Carrito.objects.create(usuario=self.user)
        ItemCarrito.objects.create(
            carrito=carrito, memoria_ram=self.otro, cantidad=1, precio_unitario=self.otro.precio,
        )
        respuesta = self.post(reverse('core:agregar_armado_al_carrito'), [
            {'tipo': 'procesador', 'id': self.cpu.id},
            {'tipo': 'memoria_ram', 'id': self.otro.id, 'cantidad': 2},
            {'tipo': 'memoria_ram', 'id': self.modulo.id, 'cantidad': 1},
        ])

        self.assertEqual(respuesta.status_code, 400)
        self.assertEqual(ItemCarrito.objects.count(), 1)
        self.assertEqual(ItemCarrito.objects.get().cantidad, 1)
        self.assertFalse(ItemCarrito.objects.filter(procesador__isnull=False).exists())

    def test_un_rechazo_de_compatibilidad_no_agrega_parcialmente(self):
        respuesta = self.post(reverse('core:agregar_armado_al_carrito'), [
            {'tipo': 'procesador', 'id': self.cpu.id},
            {'tipo': 'placa_madre', 'id': self.placa.id},
            {'tipo': 'memoria_ram', 'id': self.modulo.id, 'cantidad': 1},
            {'tipo': 'memoria_ram', 'id': self.sodimm.id, 'cantidad': 1},
        ])

        self.assertEqual(respuesta.status_code, 400)
        self.assertEqual(respuesta.json()['estado'], 'incompatible')
        self.assertEqual(ItemCarrito.objects.count(), 0)

    def test_datos_desconocidos_bloquean_y_la_compra_individual_no(self):
        armado = self.post(reverse('core:agregar_armado_al_carrito'), [
            {'tipo': 'placa_madre', 'id': self.placa.id},
            {'tipo': 'memoria_ram', 'id': self.sin_dato.id, 'cantidad': 1},
        ])
        self.assertEqual(armado.status_code, 400)
        self.assertEqual(armado.json()['estado'], 'datos_insuficientes')
        self.assertEqual(ItemCarrito.objects.count(), 0)

        individual = self.client.post(reverse('core:agregar_al_carrito'), {
            'product_id': self.sin_dato.id,
            'model_name': 'memoria_ram',
            'quantity': 1,
        })
        self.assertEqual(individual.status_code, 302)
        self.assertEqual(ItemCarrito.objects.get().memoria_ram_id, self.sin_dato.id)

    def test_cantidades_validas_se_suman_en_el_carrito(self):
        respuesta = self.post(reverse('core:agregar_armado_al_carrito'), [
            {'tipo': 'placa_madre', 'id': self.placa.id},
            {'tipo': 'memoria_ram', 'id': self.modulo.id, 'cantidad': 1},
            {'tipo': 'memoria_ram', 'id': self.modulo.id, 'cantidad': 1},
        ])

        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(respuesta.json()['estado'], 'incompleto')
        item = ItemCarrito.objects.get(memoria_ram__isnull=False)
        self.assertEqual(item.cantidad, 2)
        self.assertEqual(item.memoria_ram_id, self.modulo.id)
