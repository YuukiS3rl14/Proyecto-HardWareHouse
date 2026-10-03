import json
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from core.compatibilidad import MENSAJE_ALMACENAMIENTO
from core.excel_armado import ANCHOS, construir_libro
from core.models import AlmacenamientoSSD, Carrito, Gabinete, ItemCarrito, Proveedor, TarjetaGrafica


NOMBRE_GPU = (
    'Tarjeta gráfica de demostración con un nombre muy largo para comprobar '
    'que el ajuste de texto muestra la referencia completa en el presupuesto'
)
NOMBRE_CASE = (
    'Gabinete de torre completa para prueba de exportación con descripción '
    'extensa del espacio interno y del límite de largo de la tarjeta gráfica'
)


def _fila(componente, producto, precio, estado, texto, especificaciones):
    return {
        'componente': componente,
        'producto': producto,
        'cantidad': 1,
        'precio': precio,
        'estado': estado,
        'texto_estado': texto,
        'especificaciones': especificaciones,
    }


class ExcelArmadoTests(TestCase):
    def test_el_libro_se_lee_completo_con_el_conflicto_de_largo(self):
        motivo = (
            'La GPU mide 321 mm y el gabinete admite hasta 320 mm. '
            'Solo se comprueba el largo; no el grosor, la altura ni el espacio de radiadores.'
        )
        filas = [
            _fila(
                'Tarjeta Gráfica (GPU)',
                NOMBRE_GPU,
                219990,
                'incompatible',
                f'Incompatible\n{motivo}',
                [('Largo GPU (mm)', '321'), ('VRAM (GB)', '8')],
            ),
            _fila(
                'Gabinete',
                NOMBRE_CASE,
                59990,
                'incompatible',
                f'Incompatible\n{motivo}',
                [('Largo máximo de GPU (mm)', '320'), ('Formato admitido', 'ATX')],
            ),
            _fila(
                'Almacenamiento',
                'Unidad de estado sólido de prueba con nombre largo para la hoja de especificaciones',
                45990,
                'no_evaluada',
                'Compatibilidad no evaluada',
                [('Capacidad (GB)', '1000')],
            ),
        ]
        resumen = {
            'estado': 'incompatible',
            'texto': 'Incompatible (Revisar Componentes)',
            'motivos': [motivo],
            'aclaraciones': [
                'Almacenamiento: Compatibilidad no evaluada. Esta pieza no tiene una regla de compatibilidad.',
            ],
        }
        fecha = timezone.make_aware(datetime(2026, 10, 3, 18, 30))
        libro = construir_libro(filas, resumen, fecha)
        presupuesto = libro['Presupuesto']
        especificaciones = libro['Especificaciones']

        self.assertEqual(presupuesto['A1'].value, 'Presupuesto de armado de PC')
        self.assertIn('3 de octubre de 2026', presupuesto['A2'].value)
        self.assertIn('18:30 (Hora de Chile)', presupuesto['A2'].value)
        self.assertEqual(presupuesto['A2'].font.size, 12)
        self.assertTrue(presupuesto['A2'].font.bold)
        self.assertEqual(presupuesto['A2'].font.color.rgb[-6:], '1B4332')
        self.assertGreaterEqual(presupuesto.row_dimensions[2].height, 24)
        self.assertIn('Hora de Chile', especificaciones['A2'].value)
        self.assertEqual(especificaciones['A2'].font.size, 12)
        self.assertTrue(especificaciones['A2'].font.bold)
        self.assertEqual(
            [presupuesto.cell(4, columna).value for columna in range(1, 7)],
            ['Componente', 'Producto', 'Cantidad', 'Precio unitario', 'Subtotal', 'Estado'],
        )
        self.assertEqual(presupuesto['B5'].value, NOMBRE_GPU)
        self.assertEqual(presupuesto['D5'].value, 219990)
        self.assertEqual(presupuesto['E5'].value, 219990)
        self.assertEqual(presupuesto['D6'].value, 59990)
        self.assertEqual(presupuesto['E6'].value, 59990)
        self.assertEqual(presupuesto['E7'].value, 45990)
        self.assertEqual(presupuesto['E8'].value, 219990 + 59990 + 45990)
        self.assertEqual(presupuesto['D5'].number_format, '"$"#,##0')
        self.assertEqual(presupuesto['E8'].number_format, '"$"#,##0')
        self.assertIn('321 mm', presupuesto['F5'].value)
        self.assertIn('320 mm', presupuesto['F5'].value)
        self.assertGreater(presupuesto.row_dimensions[5].height, 30)
        self.assertTrue(presupuesto['B5'].alignment.wrap_text)
        self.assertTrue(presupuesto['F5'].alignment.wrap_text)
        self.assertGreaterEqual(presupuesto.column_dimensions['B'].width, 40)
        self.assertGreaterEqual(presupuesto.column_dimensions['F'].width, 60)
        self.assertEqual(presupuesto.freeze_panes, 'A5')
        self.assertEqual(presupuesto.page_setup.orientation, 'landscape')
        self.assertEqual(presupuesto.page_setup.fitToWidth, 1)
        self.assertTrue(presupuesto.sheet_properties.pageSetUpPr.fitToPage)
        self.assertEqual(presupuesto['F5'].fill.fgColor.rgb[-6:], 'F8D7DA')
        self.assertEqual(presupuesto['F7'].fill.fgColor.rgb[-6:], 'E9ECEF')
        self.assertEqual(presupuesto['A8'].fill.fgColor.rgb[-6:], '1B4332')
        textos = ' '.join(
            str(valor) for fila in presupuesto.iter_rows(min_row=9, max_col=1, values_only=True) for valor in fila if valor
        )
        self.assertIn(motivo, textos)
        self.assertTrue(any(
            valor == 'Largo GPU (mm)'
            for fila in especificaciones.iter_rows(min_row=5, min_col=3, max_col=3, values_only=True)
            for valor in fila
        ))
        self.assertTrue(any(
            valor == 'Largo máximo de GPU (mm)'
            for fila in especificaciones.iter_rows(min_row=5, min_col=3, max_col=3, values_only=True)
            for valor in fila
        ))
        self.assertNotIn('largo_mm', ' '.join(
            str(valor) for fila in especificaciones.iter_rows(values_only=True) for valor in fila if valor
        ))
        self.assertEqual(especificaciones.freeze_panes, 'A5')
        self.assertEqual(especificaciones.page_setup.orientation, 'landscape')

    def test_los_demas_estados_se_distinguen(self):
        libro = construir_libro(
            [
                _fila('Procesador (CPU)', 'CPU', 1000, 'compatible', 'Coincide: Socket AM5', []),
                _fila('Memoria RAM', 'RAM', 2000, 'datos_insuficientes', 'Datos insuficientes', []),
                _fila('Gabinete', 'Caja', 3000, 'incompleto', 'Selección incompleta', []),
            ],
            {
                'estado': 'datos_insuficientes',
                'texto': 'Datos insuficientes (no se declara compatible)',
                'motivos': ['No se puede comprobar el largo porque falta el largo de la GPU.'],
                'aclaraciones': [],
            },
            timezone.now(),
        )
        hoja = libro['Presupuesto']
        self.assertEqual(hoja['F5'].fill.fgColor.rgb[-6:], 'D1E7DD')
        self.assertEqual(hoja['F6'].fill.fgColor.rgb[-6:], 'FFF3CD')
        self.assertEqual(hoja['F7'].fill.fgColor.rgb[-6:], 'E7F1FF')
        self.assertEqual(hoja['E8'].value, 6000)

    def test_la_aclaracion_larga_usa_el_ancho_combinado(self):
        aclaracion = f'Almacenamiento: {MENSAJE_ALMACENAMIENTO}'
        motivo = (
            'La GPU mide 321 mm y el gabinete admite hasta 320 mm. '
            'Solo se comprueba el largo; no el grosor, la altura ni el espacio de radiadores. '
            'El texto sigue para comprobar que un motivo largo no queda recortado en la fila.'
        )
        libro = construir_libro(
            [
                _fila(
                    'Almacenamiento',
                    'SSD',
                    10000,
                    'no_evaluada',
                    MENSAJE_ALMACENAMIENTO,
                    [('Capacidad (GB)', '1000')],
                ),
            ],
            {
                'estado': 'incompleto',
                'texto': 'Selección incompleta',
                'motivos': [motivo],
                'aclaraciones': [aclaracion],
            },
            timezone.now(),
        )
        hoja = libro['Presupuesto']
        especificaciones = libro['Especificaciones']
        self.assertGreater(len(aclaracion), int(sum(ANCHOS) * 0.95))
        fila_aclaracion = next(
            numero for numero in range(1, 25) if hoja.cell(numero, 1).value == aclaracion
        )
        fila_motivo = next(
            numero for numero in range(1, 25) if hoja.cell(numero, 1).value == motivo
        )
        self.assertGreaterEqual(hoja.row_dimensions[fila_aclaracion].height, 34)
        self.assertGreaterEqual(hoja.row_dimensions[fila_motivo].height, 34)
        self.assertGreaterEqual(hoja.row_dimensions[5].height, 34)
        self.assertTrue(hoja.cell(fila_aclaracion, 1).alignment.wrap_text)
        self.assertGreaterEqual(especificaciones.row_dimensions[5].height, 18)


class ExportarArmadoTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        proveedor = Proveedor.objects.create(nombre='Marca Excel')
        cls.gpu = TarjetaGrafica.objects.create(
            proveedor=proveedor, nombre=NOMBRE_GPU, precio=Decimal('219990.00'), stock=2,
            vram_gb=8, tipo_memoria='GDDR6', interfaz='PCIe 4.0', largo_mm=321,
        )
        cls.case = Gabinete.objects.create(
            proveedor=proveedor, nombre=NOMBRE_CASE, precio=Decimal('59990.00'), stock=2,
            formato_soporte='ATX', material='Acero', largo_max_gpu_mm=320,
        )
        cls.ssd = AlmacenamientoSSD.objects.create(
            proveedor=proveedor, nombre='SSD de prueba', precio=Decimal('45990.00'), stock=2,
            capacidad_gb=1000, interfaz='NVMe PCIe 4.0', formato='M.2 2280',
        )
        cls.user = User.objects.create_user('excel', 'excel@example.com', 'clave-excel')

    def test_la_exportacion_usa_el_precio_guardado_y_no_toca_datos(self):
        precio_gpu = self.gpu.precio
        response = self.client.post(
            reverse('core:exportar_armado'),
            data=json.dumps({'componentes': [
                {'tipo': 'tarjeta_grafica', 'id': self.gpu.id, 'precio': '1'},
                {'tipo': 'gabinete', 'id': self.case.id, 'precio': '1'},
                {'tipo': 'almacenamiento_ssd', 'id': self.ssd.id},
            ]}),
            content_type='application/json',
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn('spreadsheetml', response['Content-Type'])
        libro = load_workbook(BytesIO(response.content))
        presupuesto = libro['Presupuesto']
        por_componente = {
            presupuesto.cell(fila, 1).value: fila
            for fila in range(5, 8)
        }
        fila_gpu = por_componente['Tarjeta Gráfica (GPU)']
        fila_case = por_componente['Gabinete']
        fila_ssd = por_componente['Almacenamiento']
        self.assertEqual(presupuesto.cell(fila_gpu, 4).value, 219990)
        self.assertEqual(presupuesto.cell(fila_case, 4).value, 59990)
        self.assertEqual(presupuesto.cell(fila_ssd, 4).value, 45990)
        self.assertEqual(presupuesto['E8'].value, 325970)
        self.assertIn('321 mm', presupuesto.cell(fila_gpu, 6).value)
        self.assertIn('320 mm', presupuesto.cell(fila_case, 6).value)
        self.assertIn('Compatibilidad no evaluada', presupuesto.cell(fila_ssd, 6).value)
        textos = ' '.join(
            str(valor) for fila in presupuesto.iter_rows(values_only=True) for valor in fila if valor
        )
        self.assertIn('radiadores', textos)
        self.assertIn('puertos SATA', textos)
        self.assertIn('ranuras M.2', textos)
        self.assertIn('bahías', textos)
        self.gpu.refresh_from_db()
        self.case.refresh_from_db()
        self.assertEqual(self.gpu.precio, precio_gpu)
        self.assertEqual(self.gpu.largo_mm, 321)
        self.assertEqual(self.case.largo_max_gpu_mm, 320)
        self.assertEqual(Carrito.objects.count(), 0)
        self.assertEqual(ItemCarrito.objects.count(), 0)
        self.assertFalse(User.objects.filter(username='excel-nuevo').exists())

    def test_un_error_no_devuelve_un_libro(self):
        response = self.client.post(
            reverse('core:exportar_armado'),
            data=json.dumps({'componentes': [{'tipo': 'gabinete', 'id': 999999}]}),
            content_type='application/json',
        )
        self.assertEqual(response.status_code, 400)
        self.assertNotIn('spreadsheetml', response['Content-Type'])

    def test_el_armador_pide_el_archivo_al_servidor(self):
        js = Path(__file__).resolve().parent.joinpath('static', 'core', 'js', 'pc-builder.js').read_text(encoding='utf-8')
        inicio = js.index('async function downloadAsExcel')
        cuerpo = js[inicio:js.index('URL.createObjectURL')]
        self.assertIn('exportarUrl', cuerpo)
        self.assertNotIn('evaluacionActual', cuerpo)
        self.assertNotIn('XLSX', js)
