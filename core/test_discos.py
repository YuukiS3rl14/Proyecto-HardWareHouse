import json
from decimal import Decimal
from io import BytesIO

from django.contrib.auth.models import User
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from core.compatibilidad import MENSAJE_ELEGIR_GABINETE, MENSAJE_ELEGIR_GPU, PoliticaPotencia
from core.models import (
    AlmacenamientoHDD,
    AlmacenamientoSSD,
    Carrito,
    FuenteDePoder,
    Gabinete,
    ItemCarrito,
    Procesador,
    Proveedor,
    TarjetaGrafica,
)


class OrdenYDiscosTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.proveedor = Proveedor.objects.create(nombre='Marca discos')

        cls.gpu = TarjetaGrafica.objects.create(
            proveedor=cls.proveedor, nombre='GPU 280', precio=Decimal('200000'), stock=3,
            vram_gb=8, tipo_memoria='GDDR6', interfaz='PCIe 4.0',
            largo_mm=280, consumo_referencia_watts=220,
        )
        cls.gpu_sin_largo = TarjetaGrafica.objects.create(
            proveedor=cls.proveedor, nombre='GPU sin largo', precio=Decimal('100000'), stock=3,
            vram_gb=4, tipo_memoria='GDDR6', interfaz='PCIe 4.0',
        )
        cls.mini = Gabinete.objects.create(
            proveedor=cls.proveedor, nombre='Gabinete corto', precio=Decimal('40000'), stock=3,
            formato_soporte='Mini-ITX', material='Acero', largo_max_gpu_mm=200,
        )
        cls.atx = Gabinete.objects.create(
            proveedor=cls.proveedor, nombre='Gabinete holgado', precio=Decimal('50000'), stock=3,
            formato_soporte='ATX', material='Acero', largo_max_gpu_mm=320,
        )
        cls.cpu = Procesador.objects.create(
            proveedor=cls.proveedor, nombre='CPU 65', precio=Decimal('90000'), stock=3,
            socket='AM5', nucleos=6, frecuencia_base=Decimal('3.50'),
            potencia_referencia_watts=65,
        )
        cls.fuente = FuenteDePoder.objects.create(
            proveedor=cls.proveedor, nombre='Fuente 650', precio=Decimal('70000'), stock=3,
            potencia_watts=650, certificacion='80+ Bronze',
        )
        cls.ssd = AlmacenamientoSSD.objects.create(
            id=42, proveedor=cls.proveedor, nombre='SSD compartido', precio=Decimal('10000'), stock=5,
            capacidad_gb=1000, interfaz='NVMe PCIe 4.0', formato='M.2 2280',
        )
        cls.hdd = AlmacenamientoHDD.objects.create(
            id=42, proveedor=cls.proveedor, nombre='HDD compartido', precio=Decimal('5000'), stock=5,
            capacidad_gb=2000, velocidad_rpm=7200, cache_mb=64,
        )
        cls.ssd_corto = AlmacenamientoSSD.objects.create(
            proveedor=cls.proveedor, nombre='SSD escaso', precio=Decimal('8000'), stock=1,
            capacidad_gb=500, interfaz='SATA', formato='2.5',
        )
        cls.user = User.objects.create_user('discos', 'discos@example.com', 'clave-discos')

    def setUp(self):
        self.client.login(username='discos', password='clave-discos')

    def post(self, url, payload):
        return self.client.post(url, data=json.dumps(payload), content_type='application/json')

    def test_elegir_la_gpu_primero_y_cambiar_el_gabinete(self):
        solo_gpu = self.post(reverse('core:evaluar_armado'), {
            'componentes': [{'tipo': 'tarjeta_grafica', 'id': self.gpu.id}],
        }).json()
        self.assertEqual(solo_gpu['estado'], 'incompleto')
        self.assertEqual(solo_gpu['piezas']['tarjeta_grafica']['estado'], 'incompleto')
        self.assertIn(MENSAJE_ELEGIR_GABINETE, solo_gpu['piezas']['tarjeta_grafica']['pendientes'])
        self.assertNotEqual(solo_gpu['piezas']['tarjeta_grafica']['estado'], 'datos_insuficientes')

        recomendados = self.post(reverse('core:recomendar_armado'), {
            'categoria': 'gabinete',
            'componentes': [{'tipo': 'tarjeta_grafica', 'id': self.gpu.id}],
        }).json()['candidatos']
        por_nombre = {candidato['nombre']: candidato for candidato in recomendados}
        self.assertEqual(por_nombre['Gabinete corto']['estado'], 'incompatible')
        self.assertIn('280 mm', ' '.join(por_nombre['Gabinete corto']['motivos']))
        self.assertIn('200 mm', ' '.join(por_nombre['Gabinete corto']['motivos']))
        self.assertEqual(por_nombre['Gabinete holgado']['estado'], 'compatible')
        self.assertTrue(any('280 mm' in texto for texto in por_nombre['Gabinete holgado']['coincidencias']))

        con_corto = self.post(reverse('core:evaluar_armado'), {
            'componentes': [
                {'tipo': 'tarjeta_grafica', 'id': self.gpu.id},
                {'tipo': 'gabinete', 'id': self.mini.id},
            ],
        }).json()
        self.assertEqual(con_corto['estado'], 'incompatible')
        self.assertEqual(con_corto['piezas']['gabinete']['estado'], 'incompatible')

        con_holgado = self.post(reverse('core:evaluar_armado'), {
            'componentes': [
                {'tipo': 'tarjeta_grafica', 'id': self.gpu.id},
                {'tipo': 'gabinete', 'id': self.atx.id},
            ],
        }).json()
        self.assertNotEqual(con_holgado['piezas']['gabinete']['estado'], 'incompatible')
        self.assertTrue(any('320 mm' in texto for texto in con_holgado['piezas']['gabinete']['coincidencias']))

        sin_gpu = self.post(reverse('core:evaluar_armado'), {
            'componentes': [{'tipo': 'gabinete', 'id': self.atx.id}],
        }).json()
        self.assertEqual(sin_gpu['estado'], 'incompleto')
        self.assertIn(MENSAJE_ELEGIR_GPU, sin_gpu['piezas']['gabinete']['pendientes'])

        sin_dato = self.post(reverse('core:evaluar_armado'), {
            'componentes': [
                {'tipo': 'tarjeta_grafica', 'id': self.gpu_sin_largo.id},
                {'tipo': 'gabinete', 'id': self.atx.id},
            ],
        }).json()
        self.assertEqual(sin_dato['estado'], 'datos_insuficientes')
        self.assertEqual(sin_dato['piezas']['tarjeta_grafica']['estado'], 'datos_insuficientes')
        self.assertNotEqual(sin_dato['piezas']['tarjeta_grafica']['estado'], 'incompleto')

    def test_mezcla_ssd_y_hdd_con_el_mismo_id_y_varias_unidades(self):
        from core.models import ItemCarrito

        self.assertEqual(self.ssd.id, self.hdd.id)
        response = self.post(reverse('core:agregar_armado_al_carrito'), {
            'componentes': [
                {'tipo': 'almacenamiento_ssd', 'id': self.ssd.id, 'cantidad': 2},
                {'tipo': 'almacenamiento_ssd', 'id': self.ssd.id, 'cantidad': 1},
                {'tipo': 'almacenamiento_hdd', 'id': self.hdd.id, 'cantidad': 3},
            ],
        })

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['estado'], 'compatible')
        items = list(ItemCarrito.objects.order_by('id'))
        self.assertEqual(len(items), 2)
        ssd = ItemCarrito.objects.get(almacenamiento_ssd_id=42)
        hdd = ItemCarrito.objects.get(almacenamiento_hdd_id=42)
        self.assertEqual(ssd.cantidad, 3)
        self.assertEqual(hdd.cantidad, 3)
        self.assertIsNone(ssd.almacenamiento_hdd_id)
        self.assertIsNone(hdd.almacenamiento_ssd_id)
        self.assertNotEqual(ssd.almacenamiento_ssd.nombre, hdd.almacenamiento_hdd.nombre)

    def test_cantidades_invalidas_o_sin_stock_no_agregan_nada(self):
        from core.models import Carrito, ItemCarrito

        for cantidad in (0, -1, 1.5, True, 'dos'):
            response = self.post(reverse('core:agregar_armado_al_carrito'), {
                'componentes': [
                    {'tipo': 'procesador', 'id': self.cpu.id},
                    {'tipo': 'almacenamiento_ssd', 'id': self.ssd.id, 'cantidad': cantidad},
                ],
            })
            self.assertEqual(response.status_code, 400, cantidad)
            self.assertEqual(ItemCarrito.objects.count(), 0)

        insuficiente = self.post(reverse('core:agregar_armado_al_carrito'), {
            'componentes': [
                {'tipo': 'procesador', 'id': self.cpu.id},
                {'tipo': 'almacenamiento_ssd', 'id': self.ssd_corto.id, 'cantidad': 2},
            ],
        })
        self.assertEqual(insuficiente.status_code, 400)
        self.assertIn('stock', ' '.join(insuficiente.json()['motivos']))
        self.assertEqual(ItemCarrito.objects.count(), 0)

        carrito = Carrito.objects.create(usuario=self.user)
        ItemCarrito.objects.create(
            carrito=carrito, almacenamiento_ssd=self.ssd_corto, cantidad=1,
            precio_unitario=self.ssd_corto.precio,
        )
        acumulado = self.post(reverse('core:agregar_armado_al_carrito'), {
            'componentes': [
                {'tipo': 'procesador', 'id': self.cpu.id},
                {'tipo': 'almacenamiento_ssd', 'id': self.ssd_corto.id, 'cantidad': 1},
            ],
        })
        self.assertEqual(acumulado.status_code, 400)
        self.assertEqual(ItemCarrito.objects.count(), 1)
        self.assertEqual(ItemCarrito.objects.get().cantidad, 1)
        self.assertFalse(ItemCarrito.objects.filter(procesador__isnull=False).exists())

    def test_la_compra_individual_sigue_sin_validar_el_armado(self):
        from core.models import ItemCarrito

        response = self.client.post(reverse('core:agregar_al_carrito'), {
            'product_id': self.ssd.id,
            'model_name': 'almacenamiento_ssd',
            'quantity': 2,
        })
        self.assertEqual(response.status_code, 302)
        item = ItemCarrito.objects.get()
        self.assertEqual(item.cantidad, 2)
        self.assertEqual(item.almacenamiento_ssd_id, self.ssd.id)

    def test_la_reserva_de_potencia_no_cambia_con_los_discos(self):
        self.assertEqual(PoliticaPotencia().reserva_otros_watts, 150)
        base = [
            {'tipo': 'procesador', 'id': self.cpu.id},
            {'tipo': 'tarjeta_grafica', 'id': self.gpu.id},
            {'tipo': 'fuente_de_poder', 'id': self.fuente.id},
        ]
        con_discos = base + [
            {'tipo': 'almacenamiento_ssd', 'id': self.ssd.id, 'cantidad': 2},
            {'tipo': 'almacenamiento_hdd', 'id': self.hdd.id, 'cantidad': 3},
        ]
        sin_discos = self.post(reverse('core:evaluar_armado'), {'componentes': base}).json()
        con = self.post(reverse('core:evaluar_armado'), {'componentes': con_discos}).json()
        self.assertIn('522 W', sin_discos['piezas']['fuente_de_poder']['etiqueta'])
        self.assertEqual(
            sin_discos['piezas']['fuente_de_poder']['etiqueta'],
            con['piezas']['fuente_de_poder']['etiqueta'],
        )
        self.assertIn('SATA', con['piezas']['almacenamiento']['etiqueta'])
        self.assertNotEqual(con['piezas']['almacenamiento']['estado'], 'compatible')

    def test_la_exportacion_suma_los_subtotales_de_cada_disco(self):
        response = self.post(reverse('core:exportar_armado'), {
            'componentes': [
                {'tipo': 'almacenamiento_ssd', 'id': self.ssd.id, 'cantidad': 1},
                {'tipo': 'almacenamiento_ssd', 'id': self.ssd.id, 'cantidad': 1},
                {'tipo': 'almacenamiento_hdd', 'id': self.hdd.id, 'cantidad': 3},
            ],
        })
        self.assertEqual(response.status_code, 200)
        libro = load_workbook(BytesIO(response.content))
        presupuesto = libro['Presupuesto']
        especificaciones = libro['Especificaciones']
        filas = {
            presupuesto.cell(fila, 2).value: fila
            for fila in range(5, 8)
            if presupuesto.cell(fila, 2).value
        }
        self.assertEqual(presupuesto.cell(filas['SSD compartido'], 3).value, 2)
        self.assertEqual(presupuesto.cell(filas['SSD compartido'], 5).value, 20000)
        self.assertEqual(presupuesto.cell(filas['HDD compartido'], 3).value, 3)
        self.assertEqual(presupuesto.cell(filas['HDD compartido'], 5).value, 15000)
        self.assertEqual(presupuesto.cell(7, 5).value, 35000)
        self.assertIn('Hora de Chile', presupuesto['A2'].value)
        self.assertEqual(presupuesto['A2'].font.size, 12)
        self.assertTrue(presupuesto['A2'].font.bold)
        textos = ' '.join(
            str(valor)
            for fila in especificaciones.iter_rows(values_only=True)
            for valor in fila if valor
        )
        self.assertIn('SSD compartido', textos)
        self.assertIn('HDD compartido', textos)
        self.assertIn('Hora de Chile', especificaciones['A2'].value)
        self.assertEqual(especificaciones['A2'].font.size, 12)
        self.assertTrue(especificaciones['A2'].font.bold)
