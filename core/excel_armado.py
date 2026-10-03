"""Libro de presupuesto del armador.

SheetJS 0.18.5, usado antes en el navegador, no escribe estilos de celda.
openpyxl sí aplica rellenos, bordes, formato numérico, ajuste de texto,
paneles inmovilizados e impresión.
"""

from decimal import Decimal
from io import BytesIO
from zoneinfo import ZoneInfo

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

from .compatibilidad import MENSAJE_ALMACENAMIENTO

ZONA_CHILE = ZoneInfo('America/Santiago')


COLUMNAS = (
    'Componente',
    'Producto',
    'Cantidad',
    'Precio unitario',
    'Subtotal',
    'Estado',
)
ANCHOS = (26, 48, 14, 22, 20, 64)
FORMATO_PESOS = '"$"#,##0'
MESES = (
    'enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio',
    'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre',
)

_CAMPOS = {
    'procesador': (
        ('socket', 'Socket'),
        ('nucleos', 'Núcleos'),
        ('frecuencia_base', 'Frecuencia base (GHz)'),
        ('potencia_referencia_watts', 'Potencia de referencia (W)'),
    ),
    'placa_madre': (
        ('socket_cpu', 'Socket de CPU'),
        ('chipset', 'Chipset'),
        ('formato', 'Formato'),
        ('tipo_ram_soportado', 'Tipo de RAM'),
        ('ranuras_ram', 'Ranuras de RAM'),
    ),
    'memoria_ram': (
        ('tipo_ddr', 'Tipo DDR'),
        ('capacidad_gb', 'Capacidad (GB)'),
        ('velocidad_mhz', 'Velocidad (MHz)'),
    ),
    'refrigeracion_cooler': (
        ('tipo', 'Tipo de refrigeración'),
        ('socket_compatibles', 'Sockets compatibles'),
        ('tamanho_radiador_mm', 'Radiador (mm)'),
    ),
    'tarjeta_grafica': (
        ('vram_gb', 'VRAM (GB)'),
        ('tipo_memoria', 'Tipo de memoria'),
        ('interfaz', 'Interfaz'),
        ('consumo_referencia_watts', 'Consumo de referencia (W)'),
        ('potencia_minima_fuente_watts', 'Fuente mínima recomendada (W)'),
        ('largo_mm', 'Largo GPU (mm)'),
    ),
    'almacenamiento_ssd': (
        ('capacidad_gb', 'Capacidad (GB)'),
        ('interfaz', 'Interfaz'),
        ('formato', 'Formato'),
    ),
    'almacenamiento_hdd': (
        ('capacidad_gb', 'Capacidad (GB)'),
        ('velocidad_rpm', 'Velocidad (RPM)'),
        ('cache_mb', 'Caché (MB)'),
    ),
    'gabinete': (
        ('formato_soporte', 'Formato admitido'),
        ('material', 'Material'),
        ('ventiladores_incluidos', 'Ventiladores incluidos'),
        ('largo_max_gpu_mm', 'Largo máximo de GPU (mm)'),
    ),
    'fuente_de_poder': (
        ('potencia_watts', 'Potencia (W)'),
        ('certificacion', 'Certificación'),
        ('modular', 'Modular'),
    ),
}

_TIPO_POR_MODELO = {
    'procesador': 'procesador',
    'placamadre': 'placa_madre',
    'memoriaram': 'memoria_ram',
    'refrigeracioncooler': 'refrigeracion_cooler',
    'tarjetagrafica': 'tarjeta_grafica',
    'almacenamientossd': 'almacenamiento_ssd',
    'almacenamientohdd': 'almacenamiento_hdd',
    'gabinete': 'gabinete',
    'fuentedepoder': 'fuente_de_poder',
}

_RELLENO = {
    'incompatible': ('F8D7DA', '842029'),
    'datos_insuficientes': ('FFF3CD', '664D03'),
    'incompleto': ('E7F1FF', '084298'),
    'compatible': ('D1E7DD', '0F5132'),
    'no_evaluada': ('E9ECEF', '495057'),
}
_VERDE = '1B4332'
_FILA_CLARA = 'F7F9F8'
_BORDE = Border(
    left=Side(style='thin', color='D5DDD8'),
    right=Side(style='thin', color='D5DDD8'),
    top=Side(style='thin', color='D5DDD8'),
    bottom=Side(style='thin', color='D5DDD8'),
)
_CENTRADO = Alignment(horizontal='center', vertical='center', wrap_text=True)
_IZQUIERDA = Alignment(horizontal='left', vertical='center', wrap_text=True)


def pesos_enteros(valor):
    """Devuelve el precio en pesos enteros, sin redondear un monto con centavos."""
    monto = Decimal(valor)
    if monto != monto.to_integral_value():
        raise ValueError('El precio no es un monto entero en pesos.')
    return int(monto)


def especificaciones_de(producto):
    tipo = _TIPO_POR_MODELO.get(producto._meta.model_name, '')
    visibles = []
    for campo, etiqueta in _CAMPOS.get(tipo, ()):
        valor = getattr(producto, campo, None)
        if valor is None or valor == '':
            continue
        visibles.append((etiqueta, _valor_visible(valor)))
    return visibles


def texto_de_estado(evaluacion):
    partes = []
    for texto in (evaluacion.etiqueta, *evaluacion.motivos, *evaluacion.pendientes):
        if texto and texto not in partes:
            partes.append(texto)
    return '\n'.join(partes)


def construir_libro(filas, resumen, fecha):
    """Arma el libro a partir de la evaluación recién calculada."""
    libro = Workbook()
    presupuesto = libro.active
    presupuesto.title = 'Presupuesto'
    _escribir_presupuesto(presupuesto, filas, resumen, fecha)
    _escribir_especificaciones(libro.create_sheet('Especificaciones'), filas, fecha)
    return libro


def libro_en_bytes(filas, resumen, fecha):
    buffer = BytesIO()
    construir_libro(filas, resumen, fecha).save(buffer)
    return buffer.getvalue()


def _escribir_presupuesto(hoja, filas, resumen, fecha):
    ultima = get_column_letter(len(COLUMNAS))
    hoja.merge_cells(f'A1:{ultima}1')
    hoja['A1'] = 'Presupuesto de armado de PC'
    hoja['A1'].font = Font(name='Calibri', size=18, bold=True, color=_VERDE)
    hoja['A1'].alignment = _IZQUIERDA
    hoja.row_dimensions[1].height = 28

    hoja.merge_cells(f'A2:{ultima}2')
    hoja['A2'] = f'Fecha de exportación: {_fecha_legible(fecha)}'
    _estilo_fecha(hoja['A2'])
    hoja.row_dimensions[2].height = 24

    for indice, titulo in enumerate(COLUMNAS, start=1):
        celda = hoja.cell(4, indice, titulo)
        _estilo_encabezado(celda)
    hoja.row_dimensions[4].height = 22
    hoja.freeze_panes = 'A5'
    hoja.print_title_rows = '4:4'

    for desplazamiento, fila in enumerate(filas):
        numero = 5 + desplazamiento
        valores = (
            fila['componente'],
            fila['producto'],
            fila['cantidad'],
            fila['precio'],
            fila['precio'] * fila['cantidad'],
            fila['texto_estado'],
        )
        for indice, valor in enumerate(valores, start=1):
            celda = hoja.cell(numero, indice, valor)
            celda.border = _BORDE
            celda.font = Font(name='Calibri', size=11, color='1F2933')
            if indice == 6:
                _pintar_estado(celda, fila['estado'])
            else:
                celda.fill = PatternFill(
                    'solid',
                    fgColor=_FILA_CLARA if desplazamiento % 2 else 'FFFFFF',
                )
                celda.alignment = _CENTRADO if indice in (3, 4, 5) else _IZQUIERDA
            if indice in (4, 5):
                celda.number_format = FORMATO_PESOS
        hoja.row_dimensions[numero].height = _alto(
            (
                (fila['componente'], ANCHOS[0]),
                (fila['producto'], ANCHOS[1]),
                (fila['texto_estado'], ANCHOS[5]),
            )
        )

    total_fila = 5 + len(filas)
    hoja.merge_cells(start_row=total_fila, start_column=1, end_row=total_fila, end_column=4)
    total = hoja.cell(total_fila, 1, 'Total')
    total.font = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
    total.fill = PatternFill('solid', fgColor=_VERDE)
    total.alignment = Alignment(horizontal='right', vertical='center')
    total.border = _BORDE
    for columna in range(2, 5):
        celda = hoja.cell(total_fila, columna, None)
        celda.fill = PatternFill('solid', fgColor=_VERDE)
        celda.border = _BORDE
    monto = hoja.cell(total_fila, 5, sum(fila['precio'] * fila['cantidad'] for fila in filas))
    monto.number_format = FORMATO_PESOS
    monto.font = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
    monto.fill = PatternFill('solid', fgColor=_VERDE)
    monto.alignment = _CENTRADO
    monto.border = _BORDE
    estado_total = hoja.cell(total_fila, 6, resumen['texto'])
    _pintar_estado(estado_total, resumen['estado'])
    estado_total.font = Font(name='Calibri', size=11, bold=True, color=_RELLENO[resumen['estado']][1])
    estado_total.border = _BORDE
    hoja.row_dimensions[total_fila].height = _alto(((resumen['texto'], ANCHOS[5]),))

    _escribir_validacion(hoja, total_fila + 2, resumen)
    _preparar_hoja(hoja)


def _escribir_validacion(hoja, fila, resumen):
    ultima = len(COLUMNAS)
    hoja.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ultima)
    titulo = hoja.cell(fila, 1, 'Validación del armado')
    _estilo_encabezado(titulo)
    hoja.row_dimensions[fila].height = 22

    fila += 1
    hoja.cell(fila, 1, 'Estado general').font = Font(name='Calibri', size=11, bold=True)
    hoja.merge_cells(start_row=fila, start_column=2, end_row=fila, end_column=ultima)
    estado = hoja.cell(fila, 2, resumen['texto'])
    _pintar_estado(estado, resumen['estado'])
    estado.font = Font(name='Calibri', size=11, bold=True, color=_RELLENO[resumen['estado']][1])
    hoja.row_dimensions[fila].height = _alto(((resumen['texto'], sum(ANCHOS) - ANCHOS[0]),))

    fila += 2
    hoja.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ultima)
    motivos_titulo = hoja.cell(fila, 1, 'Motivos')
    _estilo_encabezado(motivos_titulo)
    hoja.row_dimensions[fila].height = 22

    motivos = list(resumen['motivos'])
    if not motivos:
        motivos = ['No hay motivos de rechazo.']
    for motivo in motivos:
        fila += 1
        hoja.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ultima)
        celda = hoja.cell(fila, 1, motivo)
        celda.alignment = _IZQUIERDA
        celda.font = Font(name='Calibri', size=11, color='1F2933')
        if resumen['motivos']:
            _pintar_estado(celda, resumen['estado'])
        hoja.row_dimensions[fila].height = _alto(((motivo, sum(ANCHOS)),))

    aclaraciones = list(resumen['aclaraciones'])
    if aclaraciones:
        fila += 2
        hoja.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ultima)
        titulo_aclaracion = hoja.cell(fila, 1, 'Aclaraciones')
        _estilo_encabezado(titulo_aclaracion)
        hoja.row_dimensions[fila].height = 22
        for aclaracion in aclaraciones:
            fila += 1
            hoja.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ultima)
            celda = hoja.cell(fila, 1, aclaracion)
            celda.alignment = _IZQUIERDA
            celda.font = Font(name='Calibri', size=11, color='495057')
            celda.fill = PatternFill('solid', fgColor=_RELLENO['no_evaluada'][0])
            hoja.row_dimensions[fila].height = _alto(((aclaracion, sum(ANCHOS)),))


def _escribir_especificaciones(hoja, filas, fecha):
    encabezados = ('Componente', 'Producto', 'Especificación', 'Valor')
    anchos = (26, 48, 38, 42)
    hoja.merge_cells('A1:D1')
    hoja['A1'] = 'Especificaciones del armado'
    hoja['A1'].font = Font(name='Calibri', size=18, bold=True, color=_VERDE)
    hoja['A1'].alignment = _IZQUIERDA
    hoja.row_dimensions[1].height = 28
    hoja.merge_cells('A2:D2')
    hoja['A2'] = f'Fecha de exportación: {_fecha_legible(fecha)}'
    _estilo_fecha(hoja['A2'])
    hoja.row_dimensions[2].height = 24

    for indice, titulo in enumerate(encabezados, start=1):
        _estilo_encabezado(hoja.cell(4, indice, titulo))
    hoja.row_dimensions[4].height = 22
    hoja.freeze_panes = 'A5'
    hoja.print_title_rows = '4:4'

    numero = 5
    for desplazamiento, fila in enumerate(filas):
        datos = (('Cantidad', str(fila['cantidad'])),) + tuple(
            fila['especificaciones'] or (('Sin datos técnicos registrados', '—'),)
        )
        for especificacion, valor in datos:
            valores = (fila['componente'], fila['producto'], especificacion, valor)
            for indice, texto in enumerate(valores, start=1):
                celda = hoja.cell(numero, indice, texto)
                celda.border = _BORDE
                celda.font = Font(name='Calibri', size=11, color='1F2933')
                celda.alignment = _IZQUIERDA
                celda.fill = PatternFill(
                    'solid',
                    fgColor=_FILA_CLARA if desplazamiento % 2 else 'FFFFFF',
                )
            hoja.row_dimensions[numero].height = _alto(
                (
                    (fila['producto'], anchos[1]),
                    (especificacion, anchos[2]),
                    (valor, anchos[3]),
                )
            )
            numero += 1

    for indice, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(indice)].width = ancho
    _preparar_hoja(hoja)


def _preparar_hoja(hoja):
    anchos = ANCHOS if hoja.title == 'Presupuesto' else None
    if anchos:
        for indice, ancho in enumerate(anchos, start=1):
            hoja.column_dimensions[get_column_letter(indice)].width = ancho
    hoja.page_setup.orientation = 'landscape'
    hoja.page_setup.paperSize = hoja.PAPERSIZE_A4
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 0
    hoja.sheet_properties.pageSetUpPr.fitToPage = True
    hoja.page_setup.horizontalCentered = True
    hoja.page_margins = PageMargins(left=0.4, right=0.4, top=0.6, bottom=0.6, header=0.2, footer=0.2)
    hoja.sheet_view.showGridLines = False
    hoja.oddFooter.center.text = 'HardWareHouse'
    hoja.oddFooter.center.font = 'Calibri'
    hoja.oddFooter.center.size = 9


def _estilo_fecha(celda):
    celda.font = Font(name='Calibri', size=12, bold=True, color='1B4332')
    celda.alignment = _IZQUIERDA


def _estilo_encabezado(celda):
    celda.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    celda.fill = PatternFill('solid', fgColor=_VERDE)
    celda.alignment = _CENTRADO
    celda.border = _BORDE


def _pintar_estado(celda, estado):
    fondo, tinta = _RELLENO[estado]
    celda.fill = PatternFill('solid', fgColor=fondo)
    celda.font = Font(name='Calibri', size=11, color=tinta)
    celda.alignment = _IZQUIERDA


def _fecha_legible(fecha):
    if timezone.is_aware(fecha):
        fecha = fecha.astimezone(ZONA_CHILE)
    else:
        fecha = fecha.replace(tzinfo=ZONA_CHILE)
    return (
        f'{fecha.day} de {MESES[fecha.month - 1]} de {fecha.year}, '
        f'{fecha:%H:%M} (Hora de Chile)'
    )


def _valor_visible(valor):
    if isinstance(valor, bool):
        return 'Sí' if valor else 'No'
    if isinstance(valor, Decimal):
        texto = format(valor, 'f')
        if '.' in texto:
            texto = texto.rstrip('0').rstrip('.')
        return texto
    return str(valor)


def _capacidad(ancho):
    """Caracteres que caben en el ancho de Excel, con margen para el relleno."""
    return max(8, int(ancho * 0.95))


def _lineas_de_parrafo(parrafo, capacidad):
    if not parrafo:
        return 1
    lineas = 1
    usado = 0
    for token in parrafo.split(' '):
        resto = token
        while resto:
            if usado == 0:
                if len(resto) <= capacidad:
                    usado = len(resto)
                    resto = ''
                else:
                    lineas += 1
                    resto = resto[capacidad:]
            elif usado + 1 + len(resto) <= capacidad:
                usado += 1 + len(resto)
                resto = ''
            else:
                lineas += 1
                usado = 0
    return lineas


def _lineas(texto, ancho):
    """Cuenta líneas según el ancho real, incluido el de un rango combinado."""
    capacidad = _capacidad(ancho)
    total = 0
    for parte in str(texto or '').split('\n'):
        total += _lineas_de_parrafo(parte, capacidad)
    return max(1, total)


def _alto(pares):
    lineas = max(_lineas(texto, ancho) for texto, ancho in pares)
    return 15 * lineas + 4 if lineas > 1 else 18


def aclaracion_de_pieza(componente, evaluacion):
    if evaluacion.estado != 'no_evaluada':
        return None
    return f'{componente}: {MENSAJE_ALMACENAMIENTO}'
